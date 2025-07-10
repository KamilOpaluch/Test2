import os
import re
import shutil
import pandas as pd
from datetime import datetime
from collections import defaultdict
from openpyxl import load_workbook

# Define folders
base_folder = os.path.join(os.path.expanduser('~'), 'Documents', 'QE_Proxy')
output_folder = os.path.join(base_folder, 'Processed')
template_path = os.path.join(output_folder, 'Template.xlsx')
os.makedirs(output_folder, exist_ok=True)

# Regex pattern to match filenames
pattern = re.compile(r"^PROXY_(.+?)_(.+?)_(.+?)_(\d{8})\.xlsx$")

# Store grouped data
grouped_files = defaultdict(list)
unmatched_files = []

# Walk through the folder and collect files
for filename in os.listdir(base_folder):
    if filename.endswith('.xlsx') and filename.startswith('PROXY_'):
        match = pattern.match(filename)
        if match:
            name, typ, typ2, date_str = match.groups()
            group_key = (name, typ, typ2)
            file_path = os.path.join(base_folder, filename)
            grouped_files[group_key].append((file_path, date_str))

# Process each group
for (name, typ, typ2), files in grouped_files.items():
    if len(files) < 2:
        unmatched_files.append(files[0][0])
        continue

    all_data = []
    for file_path, date_str in files:
        try:
            df = pd.read_excel(file_path, skiprows=2)
            date_formatted = datetime.strptime(date_str, "%Y%m%d").strftime("%m/%d/%Y")
            df.insert(0, 'Date', date_formatted)
            all_data.append(df)
        except Exception as e:
            print(f"Error reading {file_path}: {e}")

    if all_data:
        combined_df = pd.concat(all_data, ignore_index=True)
        output_filename = f"APPENDED_{name}_{typ}_{typ2}.xlsx"
        output_path = os.path.join(output_folder, output_filename)

        # Step 1: Copy template
        if not os.path.exists(template_path):
            print(f"Template.xlsx not found at {template_path}")
            continue
        shutil.copyfile(template_path, output_path)

        # Step 2: Load workbook and write to A:L
        try:
            wb = load_workbook(output_path)
            ws = wb.active  # assuming data goes into the first sheet

            # Write headers
            for col_num, column_name in enumerate(combined_df.columns[:12], start=1):
                ws.cell(row=1, column=col_num, value=column_name)

            # Write rows
            for row_idx, row in combined_df.iterrows():
                for col_idx in range(12):  # A to L = 12 columns
                    ws.cell(row=row_idx + 2, column=col_idx + 1, value=row.iloc[col_idx])

            wb.save(output_path)
        except Exception as e:
            print(f"Error writing to {output_filename}: {e}")

# Report unmatched files
if unmatched_files:
    print("\nSkipped the following files due to no matching pair:")
    for f in unmatched_files:
        print(f"- {os.path.basename(f)}")
else:
    print("\nAll files were matched and processed.")

print("\nProcessing complete. Files saved using Template.xlsx in 'Processed' subfolder.")
